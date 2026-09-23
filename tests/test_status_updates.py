"""Regression tests for status updates after document generation."""

from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

import pandas as pd
import pytest
from docx import Document
from openpyxl import Workbook, load_workbook

from engines.docxtpl_engine import DocxtplEngine
from flow.pipeline import Pipeline
from setup import ui
from sources.excel import ExcelSource


def version(folder):
    return SimpleNamespace(
        label="test", check_records=False, date_columns=[],
        output_prefix="test", timestamp_format="%Y%m%d%H%M%S",
        output_folder=lambda: str(folder),
    )


def test_failed_render_marks_only_successful_records(tmp_path, monkeypatch, capsys):
    class Job:
        def __init__(self, record):
            self.record = record

        def result(self):
            if self.record["Serial"] == "2/2026":
                raise ValueError("render failed")
            document = Document()
            document.add_paragraph(self.record["Serial"])
            output = BytesIO()
            document.save(output)
            return output.getvalue()

    class Workers:
        def __init__(self, max_workers):
            pass

        def __enter__(self):
            return self

        def __exit__(self, *args):
            pass

        def submit(self, render, template, record):
            return Job(record)

    monkeypatch.setattr("concurrent.futures.ProcessPoolExecutor", Workers)
    questions = []
    monkeypatch.setattr(ui, "confirm", lambda question: questions.append(question) or True)
    notifications = []
    monkeypatch.setattr(ui, "notify", notifications.append)

    class Source:
        label = "test"

        def __init__(self):
            self.updated = None

        def validate_source(self):
            pass

        def load_records(self):
            return pd.DataFrame({
                "Serial": ["1/2026", "2/2026", "3/2026"],
                "Name": ["Asha", "Bhavik", "Chris"],
            })

        def mark_printed(self, data):
            self.updated = list(data["Serial"])

    source = Source()
    pipeline = Pipeline(version(tmp_path), source, DocxtplEngine("unused"))
    monkeypatch.setattr(pipeline, "validate_template_file", lambda: None)
    pipeline.run()

    assert source.updated == ["1/2026", "3/2026"]
    assert "Mark 2 as printed? (1 failed)" in questions[0]
    assert "record 2 (Serial: 2/2026, Name: Bhavik)" in capsys.readouterr().out
    assert "1 record(s) failed and were left pending" in notifications[0]


def test_excel_leaves_rows_added_mid_run_pending(tmp_path):
    path = tmp_path / "records.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Records"
    sheet.append(["Serial", "STATUS"])
    sheet.append(["1/2026", "IN PROCESS"])
    workbook.save(path)

    source = ExcelSource(str(path), "Records")
    data = source.load_records()
    workbook = load_workbook(path)
    workbook["Records"].append(["2/2026", "IN PROCESS"])
    workbook.save(path)

    source.mark_printed(data)
    sheet = load_workbook(path)["Records"]
    assert sheet["B2"].value == "PRINTED"
    assert sheet["B3"].value == "IN PROCESS"
    assert sheet["C3"].value is None


def test_locked_workbook_has_no_success_message(tmp_path, monkeypatch):
    path = tmp_path / "records.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Records"
    sheet.append(["Serial", "STATUS"])
    sheet.append(["1/2026", "IN PROCESS"])
    workbook.save(path)

    class LockedWorkbook:
        def __getitem__(self, name):
            return workbook[name]

        def save(self, path):
            raise PermissionError("workbook is locked")

    monkeypatch.setattr("sources.excel.load_workbook", lambda path: LockedWorkbook())
    monkeypatch.setattr(ui, "confirm", lambda question: True)
    messages = []
    monkeypatch.setattr(ui, "notify", messages.append)

    class Engine:
        label = "test engine"

        def generate(self, records, output_path):
            Path(output_path).write_text("saved document")
            return [1]

    source = ExcelSource(str(path), "Records")
    pipeline = Pipeline(version(tmp_path), source, Engine())
    monkeypatch.setattr(pipeline, "validate_template_file", lambda: None)
    pipeline.run()

    assert len(messages) == 1
    assert "Mission accomplished" not in messages[0]
    assert "Documents were created and saved to" in messages[0]
    assert "test_BIRTH_REGISTRATION_" in messages[0]
    assert "Excel was NOT updated: workbook is locked" in messages[0]
    assert "mark these rows as PRINTED yourself" in messages[0]
    assert len(list(tmp_path.glob("*.docx"))) == 1


def test_all_failed_records_skip_status_update(tmp_path, monkeypatch):
    class Source:
        label = "test"

        def validate_source(self):
            pass

        def load_records(self):
            return pd.DataFrame({"Serial": ["1/2026"]})

        def mark_printed(self, data):
            raise AssertionError("status update should be skipped")

    class Engine:
        label = "test engine"

        def generate(self, records, output_path):
            return []

    monkeypatch.setattr(ui, "confirm", lambda question: pytest.fail("no confirmation expected"))
    messages = []
    monkeypatch.setattr(ui, "notify", messages.append)
    pipeline = Pipeline(version(tmp_path), Source(), Engine())
    monkeypatch.setattr(pipeline, "validate_template_file", lambda: None)
    pipeline.run()

    assert len(messages) == 1
    assert "Mission accomplished" not in messages[0]


def test_excel_requires_serial_column(tmp_path):
    path = tmp_path / "records.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Records"
    sheet.append(["STATUS"])
    sheet.append(["IN PROCESS"])
    workbook.save(path)

    source = ExcelSource(str(path), "Records")
    with pytest.raises(ValueError, match="Serial column not found"):
        source.mark_printed(pd.DataFrame({"Serial": ["1/2026"]}))


def test_excel_warns_when_a_serial_is_missing(tmp_path, capsys):
    path = tmp_path / "records.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Records"
    sheet.append(["Serial", "STATUS"])
    sheet.append(["1/2026", "IN PROCESS"])
    workbook.save(path)

    source = ExcelSource(str(path), "Records")
    source.mark_printed(pd.DataFrame({"Serial": ["1/2026", "2/2026"]}))

    assert "Expected to mark 2 rows but marked 1" in capsys.readouterr().out
    assert load_workbook(path)["Records"]["B2"].value == "PRINTED"
