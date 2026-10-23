"""
Tests for case files: cases/<name>.env, one per kind of document, giving it
its own files and its own checks.

Each test runs in its own project folder (DOCUMATE_ROOT), so the real
cases/ and files/ are never read or written.
"""

import datetime
import os
import shutil

import pandas as pd
import pytest
from docx import Document
from openpyxl import Workbook, load_workbook

import main
from flow import checks
from flow.checks import CheckFailed
from setup import cases
from setup import config
from setup import ui
from versions import registry

PROJECT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


@pytest.fixture
def root(monkeypatch, tmp_path):
    """An empty project folder: cases/, files/data, files/templates."""
    monkeypatch.setenv("DOCUMATE_ROOT", str(tmp_path))
    for folder in ("cases", "files/data", "files/templates", "files/output"):
        (tmp_path / folder).mkdir(parents=True)
    return tmp_path


def write_case(root, name, text):
    path = root / "cases" / (name + ".env")
    path.write_text(text, encoding="utf-8")
    return str(path)


def frame(rows):
    """A pending sheet, as the checks see it after reading."""
    data = pd.DataFrame(rows)
    data["STATUS"] = "IN PROCESS"
    return data


# ---------------------------------------------------------------------------
# Reading a case file
# ---------------------------------------------------------------------------

def test_a_case_file_sets_the_files_and_the_rules(root):
    write_case(root, "pcc", "\n".join([
        "# comment lines are fine",
        "DOCUMATE_EXCEL_PATH=PCC_2026.xlsx",
        "DOCUMATE_ID_COLUMN=Passport No",
        "DOCUMATE_TEMPLATE=",                       # blank: still asked for
        "MUST_FILL=Passport No, Applicant Name",
        "NO_REPEATS=Passport No",
        "ALLOWED=Sex: F, M ; Grade: A, B",
        "STARTS_WITH=File No: IND/MEL/",
        "DATES=Issue Date",
    ]))

    cases.load("pcc")

    assert config.setting("DOCUMATE_EXCEL_PATH", "") == "PCC_2026.xlsx"
    assert config.id_column() == "Passport No"
    assert "DOCUMATE_TEMPLATE" not in os.environ
    assert cases.must_fill() == ["Passport No", "Applicant Name"]
    assert cases.no_repeats() == ["Passport No"]
    assert cases.allowed() == {"Sex": ["F", "M"], "Grade": ["A", "B"]}
    assert cases.starts_with() == {"File No": "IND/MEL/"}
    assert cases.dates() == ["Issue Date"]
    assert "Issue Date" in config.date_columns()


def test_a_misspelt_line_stops_rather_than_being_ignored(root):
    write_case(root, "pcc", "MUST_FIL=Name\n")
    with pytest.raises(ValueError, match="does not know: MUST_FIL"):
        cases.load("pcc")


def test_an_unknown_case_names_the_ones_there_are(root):
    write_case(root, "pcc", "")
    with pytest.raises(ValueError, match="(?s)no case called \"pc\".*Cases there: pcc"):
        cases.load("pc")


def test_allowed_needs_a_column_and_values(root):
    write_case(root, "pcc", "ALLOWED=F, M\n")
    # Caught when the case is opened, before anything is read.
    with pytest.raises(ValueError, match="like  Sex: F, M"):
        cases.load("pcc")


def test_with_no_case_there_are_no_rules():
    assert cases.named_columns() == []
    assert checks.case_rule_problems(frame([{"Name": ""}])) == []


# ---------------------------------------------------------------------------
# The rules
# ---------------------------------------------------------------------------

def test_every_broken_rule_is_listed_at_once(root):
    write_case(root, "pcc", "\n".join([
        "MUST_FILL=Passport No, Name",
        "NO_REPEATS=Passport No",
        "ALLOWED=Sex: F, M",
        "STARTS_WITH=File No: IND/",
        "DATES=Issued",
    ]))
    cases.load("pcc")
    data = frame([
        {"Passport No": "P1", "Name": "Asha", "Sex": "f", "File No": "IND/1", "Issued": "03/04/2025"},
        {"Passport No": "P1", "Name": "", "Sex": "X", "File No": "AUS/2", "Issued": "8.11.26"},
        {"Passport No": "P3", "Name": "Cy", "Sex": "", "File No": "", "Issued": None},
    ])

    found = {(row, column, issue) for row, column, issue, _value in
             checks.case_rule_problems(data)}

    assert found == {
        (3, "Name", "Missing"),
        (3, "Sex", "Not allowed - use F, M"),
        (3, "File No", "Should start with IND/"),
        (3, "Issued", "Date unreadable"),
        (2, "Passport No", "Repeated"),
        (3, "Passport No", "Repeated"),
    }


def test_numbers_are_compared_as_the_sheet_shows_them(root):
    write_case(root, "c", "ALLOWED=Year: 1, 2\nNO_REPEATS=Roll\n")
    cases.load("c")
    data = frame([{"Year": 1.0, "Roll": 7.0}, {"Year": 3.0, "Roll": 7}])
    problems = checks.case_rule_problems(data)
    assert {(r, c, v) for r, c, _i, v in problems} == {
        (3, "Year", "3"), (2, "Roll", "7"), (3, "Roll", "7")}


def test_a_rule_naming_a_column_the_sheet_lacks_stops(root):
    write_case(root, "c", "MUST_FILL=Pasport No\n")
    cases.load("c")
    with pytest.raises(CheckFailed, match="Pasport No") as failure:
        checks.case_rule_problems(frame([{"Passport No": "P1"}]))
    assert "Passport No" in failure.value.message       # the real columns are listed


def test_finished_rows_are_not_checked(root, monkeypatch):
    write_case(root, "c", "MUST_FILL=Name\n")
    cases.load("c")
    monkeypatch.setenv("DOCUMATE_ID_COLUMN", "ID")
    data = pd.DataFrame([{"ID": 1, "Name": "", "STATUS": "COMPLETED"},
                         {"ID": 2, "Name": "Bo", "STATUS": "IN PROCESS"}])
    assert list(checks.check_records(data)["ID"]) == [2]


def test_check_records_stops_with_the_table(root, monkeypatch):
    write_case(root, "c", "MUST_FILL=Name\n")
    cases.load("c")
    monkeypatch.setenv("DOCUMATE_ID_COLUMN", "ID")
    data = pd.DataFrame([{"ID": 1, "Name": "", "STATUS": None}])
    with pytest.raises(CheckFailed, match="1 row\\(s\\) break this case's rules") as failure:
        checks.check_records(data, required=["Name"], blanks_ok=lambda rows: True)
    # A MUST_FILL blank is never one to be asked about.
    assert any("Missing" in line for line in failure.value.extra_lines)


# ---------------------------------------------------------------------------
# Whole runs
# ---------------------------------------------------------------------------

def student_files(root, grade="A"):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Students"
    sheet.append(["Roll No", "Student Name", "Grade", "Exam Date"])
    sheet.append([1, "Asha Rao", grade, datetime.datetime(2026, 3, 4)])
    sheet.append([2, "Bo Lee", "B", "05/03/2026"])
    excel = root / "files" / "data" / "students.xlsx"
    workbook.save(excel)

    document = Document()
    document.add_paragraph("{{ Roll_No }} {{ Student_Name }} {{ Grade }} {{ Exam_Date }}")
    document.save(root / "files" / "templates" / "students.docx")
    return str(excel)


CASE = "\n".join([
    "DOCUMATE_EXCEL_PATH=students.xlsx",
    "DOCUMATE_TEMPLATE=students.docx",
    "DOCUMATE_ID_COLUMN=Roll No",
    "MUST_FILL=Roll No, Student Name, Grade",
    "ALLOWED=Grade: A, B, C",
    "DATES=Exam Date",
])


@pytest.fixture
def quiet_run(monkeypatch):
    """Yes to marking the rows, no PDF (that needs Word)."""
    monkeypatch.setattr(ui, "confirm", lambda question: True)
    monkeypatch.setattr(registry.VERSIONS["v3"], "to_pdf", False)


def outputs(root):
    return sorted((root / "files" / "output").glob("*.docx"))


def test_a_case_runs_end_to_end_with_no_questions(root, quiet_run):
    excel = student_files(root)
    write_case(root, "students", CASE)

    assert main.main(["v3", "students"]) == 0

    [made] = outputs(root)
    text = [p.text for p in Document(made).paragraphs if p.text.strip()]
    assert text == ["1 Asha Rao A 04/03/2026", "2 Bo Lee B 05/03/2026"]
    sheet = load_workbook(excel).active
    assert [row[4] for row in sheet.iter_rows(min_row=2, values_only=True)] == \
        ["COMPLETED", "COMPLETED"]


def test_the_case_name_alone_runs_v3(root, quiet_run):
    student_files(root)
    write_case(root, "students", CASE)
    assert main.main(["students"]) == 0
    assert len(outputs(root)) == 1


def test_a_row_breaking_the_rules_stops_the_run_and_nothing_is_made(root, quiet_run, capsys):
    excel = student_files(root, grade="Z")
    write_case(root, "students", CASE)

    main.main(["v3", "students"])

    assert outputs(root) == []
    assert "Not allowed - use A, B, C" in capsys.readouterr().out
    # Nothing was marked: the sheet never even gained a STATUS column.
    assert "STATUS" not in [c.value for c in load_workbook(excel).active[1]]


def test_the_shipped_sample_case_passes_on_the_shipped_sample(root, quiet_run):
    shutil.copy(os.path.join(PROJECT, "cases", "sample_case_study.env"), root / "cases")
    for kind, name in (("data", "SAMPLE_CASE_STUDY.xlsx"),
                       ("templates", "SAMPLE_CASE_STUDY.docx")):
        shutil.copy(os.path.join(PROJECT, "files", kind, name), root / "files" / kind)

    assert main.main(["v3", "sample_case_study"]) == 0
    assert len(outputs(root)) == 1


# ---------------------------------------------------------------------------
# How things print: REPLACE, JOIN, CAPITALS
# ---------------------------------------------------------------------------

def test_reading_the_printing_rules(root):
    write_case(root, "c", "\n".join([
        "REPLACE=Sex: F=FEMALE, m=MALE ; Title: Dr=Doctor",
        "JOIN=Address: ADD 1 + ADD 2 + ADD 3 ; Born: Date of birth + Place",
        'JOIN_SEPARATOR=" / "',
        "CAPITALS=Name, Address",
    ]))
    cases.load("c")
    assert cases.replace() == {"Sex": {"F": "FEMALE", "M": "MALE"}, "Title": {"DR": "Doctor"}}
    assert cases.joins() == [("Address", ["ADD 1", "ADD 2", "ADD 3"]),
                             ("Born", ["Date of birth", "Place"])]
    assert cases.join_separator() == " / "              # spaces kept
    assert cases.capitals() == ["Name", "Address"]
    # A JOIN's name is not a column the sheet needs.
    assert "Address" not in cases.named_columns()
    assert "ADD 2" in cases.named_columns()


def test_the_separator_defaults_to_a_comma_and_space(root):
    write_case(root, "c", "JOIN=Address: A + B\n")
    cases.load("c")
    assert cases.join_separator() == ", "


@pytest.mark.parametrize("line, message", [
    ("REPLACE=Sex: F FEMALE", "has no \"=\" in it:  F FEMALE"),
    ("JOIN=Full Address: A + B", "Call it Full_Address instead"),
    ("JOIN=Address:", "\"Address\" has none"),
])
def test_a_badly_written_printing_rule_is_named_when_the_case_opens(root, line, message):
    write_case(root, "c", line + "\n")
    with pytest.raises(ValueError, match=message):
        cases.load("c")


def test_replace_join_and_capitals_print_as_the_case_says(root):
    write_case(root, "c", "\n".join([
        "REPLACE=Sex: F=Female, M=Male",
        "JOIN=Address: ADD 1 + ADD 2 + ADD 3 ; Line: Name + Address",
        "CAPITALS=Name, Address",
    ]))
    cases.load("c")
    data = frame([
        {"Name": "asha", "Sex": "f", "ADD 1": "10 Road", "ADD 2": None, "ADD 3": " Perth "},
        {"Name": "bo", "Sex": "X", "ADD 1": "", "ADD 2": "", "ADD 3": ""},
    ])

    printed = checks.apply_case_formatting(data)

    assert list(printed["Sex"]) == ["Female", "X"]              # unlisted: as typed
    assert list(printed["Address"]) == ["10 ROAD, PERTH", ""]    # blanks left out
    assert list(printed["Line"]) == ["asha, 10 Road, Perth", "bo"]   # before CAPITALS
    assert list(printed["Name"]) == ["ASHA", "BO"]
    assert list(data["Name"]) == ["asha", "bo"]                 # the sheet is untouched


def test_with_no_printing_rules_nothing_changes():
    data = frame([{"Name": "asha"}])
    assert checks.apply_case_formatting(data) is data


def test_the_id_column_cannot_be_changed_for_printing(root, monkeypatch):
    write_case(root, "c", "CAPITALS=Passport No\n")
    cases.load("c")
    monkeypatch.setenv("DOCUMATE_ID_COLUMN", "Passport No")
    with pytest.raises(CheckFailed, match="JOIN=Passport_No_Printed: Passport No"):
        checks.case_rule_problems(frame([{"Passport No": "p1"}]))


def test_a_template_can_use_a_join_and_dates_inside_it_read_dd_mm_yyyy(root, quiet_run):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Passport No", "Name", "Sex", "Born", "Place", "ADD 1", "ADD 2"])
    sheet.append(["P1", "asha rao", "f", datetime.datetime(2020, 3, 4), "perth", "10 Road", None])
    sheet.append(["P2", "bo lee", "M", "05/06/2021", "delhi", "1 St", "Unit 2"])
    workbook.save(root / "files" / "data" / "pcc.xlsx")
    document = Document()
    document.add_paragraph("{{ Passport_No }} | {{ Name }} | {{ Sex }} | {{ Born_In }} | {{ Address }}")
    document.save(root / "files" / "templates" / "pcc.docx")
    write_case(root, "pcc", "\n".join([
        "DOCUMATE_EXCEL_PATH=pcc.xlsx",
        "DOCUMATE_TEMPLATE=pcc.docx",
        "DOCUMATE_ID_COLUMN=Passport No",
        "ALLOWED=Sex: F, M",
        "DATES=Born",
        "REPLACE=Sex: F=FEMALE, M=MALE",
        "JOIN=Born_In: Born + Place ; Address: ADD 1 + ADD 2",
        "CAPITALS=Name, Born_In, Address",
    ]))

    assert main.main(["pcc"]) == 0

    [made] = outputs(root)
    assert [p.text for p in Document(made).paragraphs if p.text.strip()] == [
        "P1 | ASHA RAO | FEMALE | 04/03/2020, PERTH | 10 ROAD",
        "P2 | BO LEE | MALE | 05/06/2021, DELHI | 1 ST, UNIT 2",
    ]
    # Marked done by the ID as typed, which the case left alone.
    statuses = [row[7] for row in load_workbook(root / "files" / "data" / "pcc.xlsx")
                .active.iter_rows(min_row=2, values_only=True)]
    assert statuses == ["COMPLETED", "COMPLETED"]
