"""
Tests for making any kind of document: whichever Excel file and Word
template are chosen at the start, with nothing set up per document.

Word and a live database are not needed. The file picker is replaced by
answers given here.
"""

import datetime
import os
import zipfile

import pandas as pd
import pytest
from docx import Document
from openpyxl import Workbook, load_workbook

import main
from engines.output_paths import build_output_name
from engines.render import render_record, with_placeholder_names
from flow import checks
from flow.checks import CheckFailed
from flow.pipeline import Pipeline
from setup import config
from setup import picker
from setup import ui
from sources import excel_format
from sources import queries_postgres
from sources import queries_tsql
from sources.database import DatabaseSource
from sources.excel import ExcelSource
from versions import registry


def student_sheet(path, rows=None, headers=None, title="Students", extra_sheet=False):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title
    sheet.append(headers or ["Roll No", "Student Name", "Marks", "Exam Date", "Remarks"])
    for row in rows or [
        [2, " Bhavik Shah ", 71, datetime.datetime(2026, 3, 5), None],
        [1, "Asha Rao", 88, datetime.datetime(2026, 3, 4), "Merit"],
    ]:
        sheet.append(row)
    if extra_sheet:
        workbook.create_sheet("Notes")
    workbook.save(path)
    return str(path)


def docx_template(path, text):
    document = Document()
    for line in text.splitlines():
        document.add_paragraph(line)
    document.save(path)
    return str(path)


def mail_merge_template(path, fields):
    """A .docx with Word Mail Merge fields, written the way Word stores them."""
    docx_template(path, "x")
    with zipfile.ZipFile(path) as original:
        parts = {name: original.read(name) for name in original.namelist()}
    body = "".join(
        '<w:p><w:fldSimple w:instr=" MERGEFIELD {} \\* MERGEFORMAT ">'
        '<w:r><w:t>x</w:t></w:r></w:fldSimple></w:p>'.format(
            '&quot;' + f + '&quot;' if " " in f else f)
        for f in fields)
    xml = parts["word/document.xml"].decode()
    parts["word/document.xml"] = xml.replace("<w:body>", "<w:body>" + body, 1).encode()
    with zipfile.ZipFile(path, "w") as rebuilt:
        for name, content in parts.items():
            rebuilt.writestr(name, content)
    return str(path)


@pytest.fixture
def students(monkeypatch):
    monkeypatch.setenv("DOCUMATE_ID_COLUMN", "Roll No")


# ---------------------------------------------------------------------------
# Nothing is assumed
# ---------------------------------------------------------------------------

def test_with_nothing_set_no_kind_of_document_is_assumed():
    assert config.data_format() == "plain"
    assert config.date_columns() == []
    assert config.sheet_name() is None


@pytest.mark.parametrize("read, name", [
    (config.excel_path, "DOCUMATE_EXCEL_PATH"),
    (lambda: config.template_path("docxtpl"), "DOCUMATE_TEMPLATE"),
    (lambda: config.template_path("mailmerge"), "DOCUMATE_TEMPLATE_MM"),
    (config.id_column, "DOCUMATE_ID_COLUMN"),
])
def test_a_run_nobody_can_be_asked_in_names_what_is_missing(read, name):
    with pytest.raises(ValueError, match=name + " is not set"):
        read()


def test_the_output_is_named_after_the_template(monkeypatch):
    monkeypatch.setenv("DOCUMATE_TEMPLATE", r"C:\anywhere\FOREIGN PCC FORMAT.docx")
    name = os.path.basename(build_output_name("out", "DocuMateX", "%d%m%Y"))
    assert name.startswith("DocuMateX_FOREIGN_PCC_FORMAT_")


# ---------------------------------------------------------------------------
# Reading the template
# ---------------------------------------------------------------------------

def test_placeholders_are_read_from_the_template(tmp_path):
    path = docx_template(tmp_path / "t.docx",
                         "{{ Passport_No }} {{ Name|upper }}\nIssued {{ year }}")
    assert picker.template_fields(path, "docxtpl") == {"Passport_No", "Name", "year"}


def test_mail_merge_fields_are_read_from_the_template(tmp_path):
    path = mail_merge_template(tmp_path / "t.docx", ["Passport_No", "Full Name"])
    assert picker.template_fields(path, "mailmerge") == {"Passport_No", "Full Name"}


def test_fields_match_headings_with_or_without_underscores():
    columns = ["Passport No", "Name", "STATUS"]
    assert picker.missing_fields({"Passport_No", "Name", "year", "Nmae"}, columns) == ["Nmae"]
    assert picker.columns_for_fields({"Passport_No", "year"}, columns) == ["Passport No"]


def test_the_likeliest_id_column_is_the_first_filled_unique_one():
    frame = pd.DataFrame({"Course": ["A", "A"], "Remarks": [None, "x"],
                          "Roll No": [1, 2], "Name": ["P", "Q"]})
    assert picker.suggest_id_column(frame) == "Roll No"


def test_the_birth_register_is_recognised_by_its_headings():
    raw = pd.DataFrame(columns=list(excel_format.RAW_COLUMNS.values()))
    prepared = pd.DataFrame(columns=excel_format.PREPARED_MARKERS)
    other = pd.DataFrame(columns=["Roll No", "Name", "Sex", "Date of birth"])
    assert excel_format.is_register_sheet(raw)
    assert excel_format.is_register_sheet(prepared)
    assert not excel_format.is_register_sheet(other)


# ---------------------------------------------------------------------------
# Choosing the files at the start
# ---------------------------------------------------------------------------

class Asked:
    """Stands in for the person at the keyboard."""

    def __init__(self, monkeypatch, files=(), confirm=False, choose=0):
        self.files = list(files)
        self.questions = []
        monkeypatch.setattr(picker, "pick_file", self.pick)
        monkeypatch.setattr(ui, "confirm", lambda q: self.questions.append(q) or confirm)
        # A number answers every list; a list of numbers answers them in turn.
        answers = list(choose) if isinstance(choose, list) else None

        def pick_from_list(prompt, options, default):
            self.questions.append(prompt)
            answer = answers.pop(0) if answers is not None else choose
            return default if answer == "default" else answer

        monkeypatch.setattr(ui, "choose", pick_from_list)

    def pick(self, title, kinds, folder):
        self.questions.append(title)
        return self.files.pop(0) if self.files else ""


def test_the_first_run_asks_for_the_files_and_the_id_column(monkeypatch, tmp_path):
    excel = student_sheet(tmp_path / "s.xlsx")
    template = docx_template(tmp_path / "t.docx", "{{ Student_Name }}")
    asked = Asked(monkeypatch, files=[excel, template], choose="default")

    picker.prepare_run(registry.VERSIONS["v3"])

    assert os.environ["DOCUMATE_EXCEL_PATH"] == excel
    assert os.environ["DOCUMATE_TEMPLATE"] == template
    assert os.environ["DOCUMATE_SHEET_NAME"] == "Students"
    assert os.environ["DOCUMATE_DATA_FORMAT"] == "plain"
    # Roll No is the first filled, unique column, so it is the default.
    assert os.environ["DOCUMATE_ID_COLUMN"] == "Roll No"
    assert len(asked.questions) == 3      # Excel, template, ID column


CHOSEN = ["DOCUMATE_EXCEL_PATH", "DOCUMATE_TEMPLATE", "DOCUMATE_SHEET_NAME",
          "DOCUMATE_ID_COLUMN", "DOCUMATE_DATA_FORMAT"]


def test_every_run_picks_the_files_again_with_last_times_as_defaults(monkeypatch, tmp_path):
    excel = student_sheet(tmp_path / "s.xlsx", extra_sheet=True)
    template = docx_template(tmp_path / "t.docx", "{{ Student_Name }}")
    Asked(monkeypatch, files=[excel, template], choose=[0, 2])   # Students; Marks as ID
    picker.prepare_run(registry.VERSIONS["v3"])
    for name in CHOSEN:
        monkeypatch.delenv(name)

    asked = Asked(monkeypatch, files=[excel, template], choose="default")
    picked_in = []
    original = asked.pick
    asked.pick = lambda title, kinds, folder: picked_in.append(folder) or \
        original(title, kinds, folder)
    monkeypatch.setattr(picker, "pick_file", asked.pick)
    picker.prepare_run(registry.VERSIONS["v3"])

    # No "use the same files?" question: Excel, template, sheet, ID column.
    assert not any("same" in str(q).lower() for q in asked.questions)
    assert len(asked.questions) == 4
    # The picker opens where last time's files were, and the lists start on
    # last time's answers.
    assert picked_in == [str(tmp_path), str(tmp_path)]
    assert os.environ["DOCUMATE_SHEET_NAME"] == "Students"
    assert os.environ["DOCUMATE_ID_COLUMN"] == "Marks"


@pytest.mark.parametrize("key", ["v3", "x", "y"])
def test_every_excel_version_picks_both_files(monkeypatch, tmp_path, key):
    excel = student_sheet(tmp_path / "s.xlsx")
    template = docx_template(tmp_path / "t.docx", "x")
    asked = Asked(monkeypatch, files=[excel, template], choose=0)
    picker.prepare_run(registry.VERSIONS[key])
    assert os.environ["DOCUMATE_EXCEL_PATH"] == excel
    assert os.environ[config.TEMPLATE_SETTINGS[registry.VERSIONS[key].engine][0]] == template


@pytest.mark.parametrize("key", ["z", "o"])
def test_the_database_versions_pick_the_template(monkeypatch, tmp_path, key):
    template = docx_template(tmp_path / "t.docx", "x")
    asked = Asked(monkeypatch, files=[template])
    picker.prepare_run(registry.VERSIONS[key])
    assert len(asked.questions) == 1
    assert os.environ[config.TEMPLATE_SETTINGS[registry.VERSIONS[key].engine][0]] == template


def test_config_env_is_read_and_blank_lines_count_as_not_set(monkeypatch, tmp_path,
                                                             read_config_env):
    (tmp_path / "config.env").write_bytes(
        b"\xef\xbb\xbfDOCUMATE_ID_COLUMN=Passport No\n"
        b"DOCUMATE_EXCEL_PATH=\nDOCUMATE_DB_PORT=\n")
    monkeypatch.setenv("DOCUMATE_ROOT", str(tmp_path))
    for name in ["DOCUMATE_ID_COLUMN", "DOCUMATE_EXCEL_PATH", "DOCUMATE_DB_PORT"]:
        monkeypatch.delenv(name, raising=False)

    try:
        assert config.id_column() == "Passport No"
        with pytest.raises(ValueError, match="DOCUMATE_EXCEL_PATH is not set"):
            config.excel_path()
        assert "DOCUMATE_DB_PORT" not in os.environ
    finally:
        os.environ.pop("DOCUMATE_ID_COLUMN", None)


def test_a_mail_merge_template_picked_for_v3_is_refused(monkeypatch, tmp_path):
    excel = student_sheet(tmp_path / "s.xlsx")
    template = mail_merge_template(tmp_path / "t_MM.docx", ["Roll_No"])
    Asked(monkeypatch, files=[excel, template], choose=0)
    with pytest.raises(ValueError, match="Mail Merge template, but v3"):
        picker.prepare_run(registry.VERSIONS["v3"])


def test_a_placeholder_template_picked_for_x_is_refused(monkeypatch, tmp_path):
    excel = student_sheet(tmp_path / "s.xlsx")
    template = docx_template(tmp_path / "t.docx", "{{ Roll_No }}")
    Asked(monkeypatch, files=[excel, template], choose=0)
    with pytest.raises(ValueError, match="placeholder }} template, but X"):
        picker.prepare_run(registry.VERSIONS["x"])


def test_choosing_nothing_stops_with_a_message(monkeypatch):
    Asked(monkeypatch, files=[])
    with pytest.raises(ValueError, match="nothing has been generated"):
        picker.prepare_run(registry.VERSIONS["v3"])


def test_files_set_in_env_are_used_without_asking(monkeypatch, tmp_path, students):
    monkeypatch.setenv("DOCUMATE_EXCEL_PATH", student_sheet(tmp_path / "s.xlsx"))
    monkeypatch.setenv("DOCUMATE_TEMPLATE", docx_template(tmp_path / "t.docx", "x"))
    asked = Asked(monkeypatch)
    picker.prepare_run(registry.VERSIONS["v3"])
    assert asked.questions == []


def test_a_mail_merge_version_asks_for_a_mail_merge_template(monkeypatch, tmp_path, students):
    monkeypatch.setenv("DOCUMATE_EXCEL_PATH", student_sheet(tmp_path / "s.xlsx"))
    template = mail_merge_template(tmp_path / "mm.docx", ["Roll_No"])
    asked = Asked(monkeypatch, files=[template])
    picker.prepare_run(registry.VERSIONS["x"])
    assert "Mail Merge" in asked.questions[0]
    assert os.environ["DOCUMATE_TEMPLATE_MM"] == template


def test_a_birth_register_sheet_needs_no_id_question(monkeypatch, tmp_path):
    headers = list(excel_format.RAW_COLUMNS.values()) + ["Sighning_Authority", "STATUS"]
    excel = student_sheet(tmp_path / "br.xlsx", headers=headers, rows=[[1] * len(headers)])
    monkeypatch.setenv("DOCUMATE_EXCEL_PATH", excel)
    monkeypatch.setenv("DOCUMATE_TEMPLATE", docx_template(tmp_path / "t.docx", "x"))
    asked = Asked(monkeypatch)

    picker.prepare_run(registry.VERSIONS["v3"])

    assert os.environ["DOCUMATE_DATA_FORMAT"] == "birth_register"
    assert config.id_column() == "Serial"
    assert asked.questions == []


def test_main_shows_why_it_stopped_when_nothing_is_chosen(monkeypatch):
    Asked(monkeypatch, files=[])
    shown = []
    monkeypatch.setattr(ui, "notify", shown.append)
    assert main.main(["v3"]) == 1
    assert "nothing has been generated" in shown[0]


# ---------------------------------------------------------------------------
# Checking the rows against the template
# ---------------------------------------------------------------------------

def pipeline_for(tmp_path, template_text, monkeypatch):
    template = docx_template(tmp_path / "t.docx", template_text)
    monkeypatch.setenv("DOCUMATE_TEMPLATE", template)
    return Pipeline(registry.VERSIONS["v3"], source=None, engine=None)


def test_a_placeholder_with_no_column_stops_the_run(monkeypatch, tmp_path, students):
    data = ExcelSource(student_sheet(tmp_path / "s.xlsx")).load_records()
    pipeline = pipeline_for(tmp_path, "{{ Student_Nmae }}", monkeypatch)
    shown = []
    monkeypatch.setattr(ui, "notify", shown.append)

    assert pipeline.check_records(data) is None
    assert "Student_Nmae" in shown[0]


def test_only_the_columns_the_template_prints_must_be_filled(monkeypatch, tmp_path, students):
    path = student_sheet(tmp_path / "s.xlsx", rows=[
        [1, "Asha", None, datetime.datetime(2026, 3, 4), None],
    ])
    data = ExcelSource(path).load_records()
    pipeline = pipeline_for(tmp_path, "{{ Student_Name }} {{ Exam_Date }}", monkeypatch)
    # Marks and Remarks are blank, but the template does not print them.
    assert len(pipeline.check_records(data)) == 1


def test_blank_printed_fields_are_asked_about(monkeypatch, tmp_path, students):
    path = student_sheet(tmp_path / "s.xlsx")    # Remarks blank on one row
    pipeline = pipeline_for(tmp_path, "{{ Student_Name }} {{ Remarks }}", monkeypatch)

    monkeypatch.setattr(ui, "confirm", lambda question: True)
    assert len(pipeline.check_records(ExcelSource(path).load_records())) == 2

    shown = []
    monkeypatch.setattr(ui, "notify", shown.append)
    monkeypatch.setattr(ui, "confirm", lambda question: False)
    assert pipeline.check_records(ExcelSource(path).load_records()) is None


def test_a_blank_id_is_never_waved_through(monkeypatch, tmp_path, students):
    data = ExcelSource(student_sheet(tmp_path / "s.xlsx")).load_records()
    data.loc[0, "Roll No"] = None
    with pytest.raises(CheckFailed):
        checks.check_records(data, required=["Student Name"],
                             blanks_ok=lambda rows: True)


def test_the_same_id_twice_is_caught(students, tmp_path):
    path = student_sheet(tmp_path / "s.xlsx", rows=[
        [1, "Asha", 88, datetime.datetime(2026, 3, 4), ""],
        [1, "Bhavik", 71, datetime.datetime(2026, 3, 5), ""],
    ])
    data = ExcelSource(path).load_records()
    with pytest.raises(CheckFailed, match="duplicate Roll No"):
        checks.check_records(data, required=["Student Name"])


# ---------------------------------------------------------------------------
# A plain sheet, read and marked
# ---------------------------------------------------------------------------

def test_a_plain_sheet_keeps_its_own_columns(students, tmp_path):
    data = ExcelSource(student_sheet(tmp_path / "s.xlsx"), "Students").load_records()
    assert list(data.columns[:5]) == ["Roll No", "Student Name", "Marks",
                                      "Exam Date", "Remarks"]
    assert list(data["Student Name"]) == ["Bhavik Shah", "Asha Rao"]
    assert list(data["STATUS"]) == [config.PENDING_STATUS] * 2


def test_whole_numbers_with_a_blank_do_not_print_as_decimals(students, tmp_path):
    path = student_sheet(tmp_path / "s.xlsx", rows=[
        [1, "Asha", 88, datetime.datetime(2026, 3, 4), "x"],
        [2, "Bhavik", None, datetime.datetime(2026, 3, 5), "y"],
    ])
    record = checks.to_records(ExcelSource(path).load_records(), add_year=False)[0]
    assert str(record["Marks"]) == "88"


def test_blank_cells_print_as_nothing(students, tmp_path):
    data = ExcelSource(student_sheet(tmp_path / "s.xlsx")).load_records()
    assert checks.to_records(data, add_year=False)[0]["Remarks"] == ""


def test_a_sheet_without_its_id_column_stops_before_anything(students, tmp_path):
    path = student_sheet(tmp_path / "s.xlsx",
                         headers=["Student Name", "Marks", "Exam Date", "Remarks", "x"])
    with pytest.raises(ValueError, match="no \"Roll No\" column"):
        ExcelSource(path).load_records()


def test_rows_are_marked_by_their_id_and_a_status_column_is_added(students, tmp_path):
    path = student_sheet(tmp_path / "s.xlsx")
    source = ExcelSource(path, "Students")
    data = source.load_records()

    source.mark_status(data[data["Roll No"] == 1], config.COMPLETED_STATUS)

    sheet = load_workbook(path)["Students"]
    headers = [cell.value for cell in sheet[1]]
    status = headers.index("STATUS") + 1
    assert sheet.cell(row=2, column=status).value is None
    assert sheet.cell(row=3, column=status).value == "COMPLETED"
    left = checks.drop_finished(ExcelSource(path, "Students").load_records())
    assert list(left["Roll No"]) == [2]


def test_with_no_sheet_named_the_first_sheet_is_used(students, tmp_path):
    path = student_sheet(tmp_path / "s.xlsx", title="Whatever")
    source = ExcelSource(path)
    source.mark_status(source.load_records(), config.COMPLETED_STATUS)
    assert "STATUS" in [cell.value for cell in load_workbook(path)["Whatever"][1]]


def test_an_id_with_stray_spaces_in_the_sheet_still_matches(students, tmp_path):
    path = student_sheet(tmp_path / "s.xlsx", rows=[
        [" A-1 ", "Asha", 88, datetime.datetime(2026, 3, 4), ""],
    ])
    source = ExcelSource(path)
    source.mark_status(source.load_records(), config.COMPLETED_STATUS)
    sheet = load_workbook(path)["Students"]
    headers = [cell.value for cell in sheet[1]]
    assert sheet.cell(row=2, column=headers.index("STATUS") + 1).value == "COMPLETED"


# ---------------------------------------------------------------------------
# Filling the template
# ---------------------------------------------------------------------------

def test_columns_with_spaces_get_a_placeholder_name():
    record = with_placeholder_names({"Student Name": "Asha", "S No.": 3})
    assert record["Student_Name"] == "Asha" and record["S_No"] == 3
    assert record["Student Name"] == "Asha"


def test_a_column_with_a_space_fills_its_placeholder(tmp_path):
    template = docx_template(tmp_path / "t.docx", "{{ Student_Name }} scored {{ Marks }}")
    output = tmp_path / "out.docx"
    output.write_bytes(render_record(template, {"Student Name": "Asha", "Marks": 88}))
    assert Document(output).paragraphs[0].text == "Asha scored 88"


# ---------------------------------------------------------------------------
# Database: a table of its own
# ---------------------------------------------------------------------------

class FakeSource(DatabaseSource):
    queries = queries_postgres

    def connect(self, settings):
        class Connection:
            def cursor(self):
                return None
        return Connection()


def test_with_no_table_named_the_birth_register_tables_are_read():
    source = FakeSource({})
    assert "FROM applicant a" in source.sql(source.queries.PENDING_APPLICANTS)
    assert "WHERE serial = ANY" in source.sql(source.queries.mark_printed_sql(1))
    assert source.row_ids(pd.DataFrame({"Serial": ["12/2026"]})) == [12]


def test_a_named_table_is_read_and_updated_by_its_own_id():
    source = FakeSource({}, schema="study", table="students", id_column="roll_no")
    assert "FROM study.students" in source.sql(source.queries.PENDING_ROWS)
    update = source.sql(source.queries.mark_printed_sql(1))
    assert "UPDATE study.students" in update and "roll_no = ANY" in update
    ids = source.row_ids(pd.DataFrame({"ROLL_NO": pd.Series([3, 4], dtype="int64")}))
    assert ids == [3, 4] and all(type(i) is int for i in ids)


def test_both_dialects_offer_the_table_query():
    assert hasattr(queries_tsql, "PENDING_ROWS")
    assert hasattr(queries_postgres, "PENDING_ROWS")


def test_a_table_name_that_is_not_a_plain_name_is_refused(monkeypatch):
    monkeypatch.setenv("DOCUMATE_DB_TABLE", "students; DROP TABLE x")
    with pytest.raises(ValueError, match="DOCUMATE_DB_TABLE"):
        config.database_table()
