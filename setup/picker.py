"""
Choosing the Excel file and Word template when DocuMate starts.

There is nothing to set up per kind of document. A run needs three things -
a workbook, a template, and the column that identifies a row - and every
version gets each the same way:

    set in config.env    used as it is, nothing asked (DOCUMATE_EXCEL_PATH, ...)
    not set        asked: a file picker for the files, a numbered list for
                   the column - every run

Everything else is worked out: which columns matter, from the template's
placeholders; which are dates, from the sheet; whether the sheet is the
birth register, from its headings.

The answers go into the environment under the setting names, so the rest of
DocuMate reads them exactly as if they had been in config.env. They are kept in
files/last_run.json only as next time's defaults: the folder the picker
opens in, and the sheet and ID column pre-selected in their lists.
"""

import json
import os
import re
import zipfile

from setup import config
from setup import messages
from setup import ui


# ---------------------------------------------------------------------------
# What the template asks for
# ---------------------------------------------------------------------------

# A Mail Merge field as Word stores it:  MERGEFIELD Passport_No \* MERGEFORMAT
# The name may be quoted when it has spaces in it.
MERGEFIELD = re.compile(r'MERGEFIELD\s+(?:"([^"]+)"|([^\s\\]+))', re.IGNORECASE)

# Word parts that can hold fields: the body, headers and footers.
WORD_PARTS = re.compile(r"word/(document|header\d*|footer\d*)\.xml$")


def template_fields(path, engine):
    """
    The field names a template uses, as a set.

    {{ placeholders }} are read by docxtpl itself, so the answer is exactly
    what it will look for. Mail Merge fields are read from the field codes
    in the file, including the ones in headers and footers.
    """
    if engine == "docxtpl":
        from docxtpl import DocxTemplate
        return set(DocxTemplate(path).get_undeclared_template_variables())

    fields = set()
    with zipfile.ZipFile(path) as document:
        for part in document.namelist():
            if not WORD_PARTS.match(part):
                continue
            xml = document.read(part).decode("utf-8", "ignore")
            # A field code can be split across several runs; joining the
            # text of every instrText first puts it back together.
            codes = "".join(re.findall(r"<w:instrText[^>]*>([^<]*)</w:instrText>", xml))
            codes += " ".join(re.findall(r'w:instr="([^"]*)"', xml))
            for quoted, plain in MERGEFIELD.findall(codes.replace("&quot;", '"')):
                fields.add(quoted or plain)
    return fields


def wrong_kind_of_template(path, engine):
    """
    True when the template was made for the other engine.

    The two kinds are easy to mix up in a file picker, and the wrong one
    does not fail - it prints documents with every value missing (or with
    {{ braces }} left in). A {{ }} template with no placeholders but with
    Mail Merge fields, or the other way round, is the giveaway.
    """
    other = "mailmerge" if engine == "docxtpl" else "docxtpl"
    own_fields = template_fields(path, engine) - {"year"}
    return not own_fields and bool(template_fields(path, other) - {"year"})


def missing_fields(fields, columns):
    """
    The template fields no column can fill, sorted.

    A field matches a column by its heading or by its placeholder name
    ("Passport No" or Passport_No). {{ year }} is always supplied.
    """
    from engines.render import placeholder_name

    available = {"year"}
    for column in columns:
        available.add(str(column))
        available.add(placeholder_name(column))
    return sorted(field for field in fields if field not in available)


def columns_for_fields(fields, columns):
    """The columns the template uses - the ones that must be filled in."""
    from engines.render import placeholder_name

    return [column for column in columns
            if str(column) in fields or placeholder_name(column) in fields]


# ---------------------------------------------------------------------------
# Remembering the last run
# ---------------------------------------------------------------------------

def memory_path():
    return os.path.join(config.files_folder(), "last_run.json")


def load_memory():
    """What was chosen last time, or {} when there is nothing usable."""
    try:
        with open(memory_path(), encoding="utf-8") as handle:
            memory = json.load(handle)
        return memory if isinstance(memory, dict) else {}
    except (OSError, ValueError):
        return {}


def save_memory(memory):
    """Remember this run's choices. Failing to is not worth stopping for."""
    try:
        os.makedirs(os.path.dirname(memory_path()), exist_ok=True)
        with open(memory_path(), "w", encoding="utf-8") as handle:
            json.dump(memory, handle, indent=2)
    except OSError:
        pass


# ---------------------------------------------------------------------------
# Asking
# ---------------------------------------------------------------------------

def pick_file(title, kinds, folder):
    """
    Ask for a file with the Windows file picker; typed in when there is none.

    Returns the full path, or "" when nothing was chosen.
    """
    try:
        import tkinter
        from tkinter import filedialog

        window = tkinter.Tk()
        window.withdraw()
        # Otherwise the picker can open behind the console window.
        window.attributes("-topmost", True)
        path = filedialog.askopenfilename(
            title=title,
            initialdir=folder if os.path.isdir(folder) else None,
            filetypes=kinds + [("All files", "*.*")],
        )
        window.destroy()
        return os.path.abspath(path) if path else ""
    except Exception:
        pass

    try:
        typed = input(messages.TYPE_PATH.format(title=title)).strip().strip('"')
    except EOFError:
        return ""
    return os.path.abspath(config.in_folder(folder, typed)) if typed else ""


EXCEL_KINDS = [("Excel workbook", "*.xlsx *.xlsm *.xls")]
WORD_KINDS = [("Word document", "*.docx")]

TEMPLATE_TITLES = {
    "docxtpl": "Choose the Word template (with {{ placeholders }})",
    "mailmerge": "Choose the Word Mail Merge template",
}


def choose_from(prompt, options, default_index):
    """ui.choose, with a closed input taking the default."""
    try:
        return ui.choose(prompt, options, default_index)
    except EOFError:
        return default_index


def suggest_id_column(frame):
    """
    The column most likely to identify a row: the first one that is filled
    in and different on every row. None when no column is.
    """
    for column in frame.columns:
        values = frame[column]
        if values.isna().any():
            continue
        text = values.astype(str).str.strip()
        if (text != "").all() and text.is_unique:
            return column
    return None


# ---------------------------------------------------------------------------
# The run
# ---------------------------------------------------------------------------

def start_folder(remembered, fallback):
    """Where the picker opens: the folder of last time's file, or `fallback`."""
    if remembered and os.path.isdir(os.path.dirname(remembered)):
        return os.path.dirname(remembered)
    return fallback


def choose_files(version, memory):
    """
    The workbook (None for a database version) and template for this run.

    Each comes from config.env when it is set there, and is picked otherwise -
    every run, every version. Last time's choice only decides which folder
    the picker opens in.
    """
    template_setting = config.TEMPLATE_SETTINGS[version.engine][0]
    needs_excel = version.source == "excel"

    excel = config.setting("DOCUMATE_EXCEL_PATH", "") if needs_excel else None
    template = config.setting(template_setting, "")

    if excel:
        excel = config.in_folder(config.data_folder(), excel)
    elif needs_excel:
        excel = pick_file("Choose the Excel file", EXCEL_KINDS,
                          start_folder(memory.get("excel"), config.data_folder()))
        if not excel:
            raise ValueError(messages.NOTHING_CHOSEN)

    if template:
        template = config.in_folder(config.templates_folder(), template)
    else:
        last = memory.get("templates", {}).get(version.engine)
        template = pick_file(TEMPLATE_TITLES[version.engine], WORD_KINDS,
                             start_folder(last, config.templates_folder()))
        if not template:
            raise ValueError(messages.NOTHING_CHOSEN)

    return excel, template


def choose_sheet(excel, memory):
    """
    The sheet to read: config.env's, the only one, or asked - with last time's
    choice for this workbook as the default.
    """
    from openpyxl import load_workbook

    named = config.sheet_name()
    if named:
        return named

    workbook = load_workbook(excel, read_only=True)
    names = workbook.sheetnames
    workbook.close()

    if len(names) == 1:
        return names[0]

    remembered = memory.get("sheets", {}).get(excel)
    default = names.index(remembered) if remembered in names else 0
    print(messages.CHOOSE_SHEET_HEADING.format(workbook=os.path.basename(excel)))
    chosen = choose_from(messages.CHOOSE_SHEET, names, default)
    if chosen is None:
        raise ValueError(messages.NOTHING_CHOSEN)
    return names[chosen]


def choose_id_column(frame, excel, memory):
    """
    The column that identifies a row: config.env's, or asked - with last time's
    answer for this workbook, or the likeliest column, as the default.
    """
    named = config.setting("DOCUMATE_ID_COLUMN", "")
    if named:
        return named

    columns = [str(column) for column in frame.columns]
    remembered = memory.get("id_columns", {}).get(excel)
    if remembered in columns:
        default = columns.index(remembered)
    else:
        suggested = suggest_id_column(frame)
        default = columns.index(str(suggested)) if suggested is not None else None

    print(messages.CHOOSE_ID_HEADING)
    chosen = choose_from(messages.CHOOSE_ID, columns, default)
    if chosen is None:
        raise ValueError(messages.NOTHING_CHOSEN)
    return columns[chosen]


def prepare_run(version):
    """
    Settle this run's files and columns: from config.env, or asked.

    Puts the answers in the environment and remembers them as next time's
    defaults. Raises ValueError with a message to show when something could
    not be settled.
    """
    import pandas as pd
    from sources import excel_format

    memory = load_memory()
    excel, template = choose_files(version, memory)

    if os.path.exists(template) and wrong_kind_of_template(template, version.engine):
        raise ValueError(messages.WRONG_TEMPLATE_KIND[version.engine].format(
            template=os.path.basename(template), version=version.label))

    os.environ[config.TEMPLATE_SETTINGS[version.engine][0]] = template
    memory.setdefault("templates", {})[version.engine] = template

    if excel:
        if not os.path.exists(excel):
            raise ValueError(messages.SOURCE_MISSING.format(path=excel))
        sheet = choose_sheet(excel, memory)
        os.environ["DOCUMATE_EXCEL_PATH"] = excel
        os.environ["DOCUMATE_SHEET_NAME"] = sheet
        memory["excel"] = excel
        memory.setdefault("sheets", {})[excel] = sheet

        frame = pd.read_excel(excel, sheet_name=sheet)

        if not config.setting("DOCUMATE_DATA_FORMAT", ""):
            register = excel_format.is_register_sheet(frame)
            os.environ["DOCUMATE_DATA_FORMAT"] = "birth_register" if register else "plain"

        if not config.is_register():
            column = choose_id_column(frame, excel, memory)
            os.environ["DOCUMATE_ID_COLUMN"] = column
            memory.setdefault("id_columns", {})[excel] = column

    save_memory(memory)

    if excel:
        print(messages.RUN_FILES.format(
            excel=os.path.basename(excel), sheet=os.environ["DOCUMATE_SHEET_NAME"],
            template=os.path.basename(template), id=config.id_column()))
    else:
        print(messages.RUN_TEMPLATE.format(template=os.path.basename(template)))
