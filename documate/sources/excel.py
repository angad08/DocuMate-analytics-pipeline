"""
The Excel data source. Shared by v3, X and Y.

Reading uses pandas. Writing back uses openpyxl on purpose - pandas would
rewrite the whole sheet and wipe out your formulas, colours and column
widths. openpyxl edits the cells and leaves the rest alone.
"""

import os
import time
from datetime import datetime

import pandas as pd
from openpyxl import load_workbook

from documate.setup import config
from documate.setup import messages


class ExcelSource:
    """Reads records from a worksheet, and writes statuses back to it."""

    # Shown in messages, e.g. "Reading Excel..."
    label = "Excel"

    def __init__(self, path, sheet):
        self.path = path
        self.sheet = sheet

    # -- reading ------------------------------------------------------------

    def validate_source(self):
        """Stop now if the file isn't there, rather than halfway through."""
        if not os.path.exists(self.path):
            raise FileNotFoundError(messages.SOURCE_MISSING.format(path=self.path))

    def load_records(self):
        """Read the sheet into a table."""
        return pd.read_excel(self.path, sheet_name=self.sheet)

    def close(self):
        """Nothing to close for a file. Here so both sources match."""
        pass

    # -- writing back -------------------------------------------------------

    def mark_printed(self, data):
        """
        Mark rows as PRINTED and stamp today's date.

        Worth knowing: this updates every row in the sheet that isn't
        already PRINTED - not only the rows we just processed. That is how
        the original worked and changing it would be a real change, not a
        tidy-up. It matters if someone adds a row to the sheet while the
        documents are being made: that row gets marked printed without ever
        getting a certificate. The database version doesn't have this
        problem because it updates by Serial number.
        """
        try:
            workbook = load_workbook(self.path)
            sheet = workbook[self.sheet]

            # Map column name -> column number, from the header row.
            columns = {}
            for number, cell in enumerate(sheet[1], start=1):
                columns[cell.value] = number

            if "STATUS" not in columns:
                print(messages.EXCEL_NO_STATUS_COLUMN)
                return

            # Older sheets don't have Date_Printed, so add it.
            if "Date_Printed" not in columns:
                new_column = sheet.max_column + 1
                sheet.cell(row=1, column=new_column, value="Date_Printed")
                columns["Date_Printed"] = new_column

            status_column = columns["STATUS"]
            date_column = columns["Date_Printed"]
            today = datetime.now().strftime("%d/%m/%Y")

            for row in range(2, sheet.max_row + 1):
                status = sheet.cell(row=row, column=status_column).value
                if status is None:
                    status = ""

                if str(status).strip().upper() == config.DONE_STATUS:
                    continue

                sheet.cell(row=row, column=status_column, value=config.DONE_STATUS)
                sheet.cell(row=row, column=date_column, value=today)

            workbook.save(self.path)

            print(messages.EXCEL_UPDATED.format(
                sheet=self.sheet,
                date=today,
                clock=time.strftime("%H:%M:%S"),
            ))

        except PermissionError:
            print(messages.EXCEL_LOCKED)

        except Exception as error:
            print(messages.EXCEL_UPDATE_ERROR.format(error=error))
